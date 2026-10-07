"""converter 模块单元测试（纯逻辑 + 临时目录真实产物，无 GUI、无真实剪贴板）。"""

from __future__ import annotations

import os
import sys

import docx
import openpyxl
import pytest

# 保证无论从哪个工作目录运行 pytest，都能 import 到项目根目录下的模块。
_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

import converter  # noqa: E402


class _FakeClipboard:
    """假剪贴板。"""

    def __init__(self, html=None, text=None) -> None:
        self._html = html
        self._text = text

    def read_html(self):
        return self._html

    def read_text(self):
        return self._text


@pytest.fixture()
def outdir(tmp_path, monkeypatch):
    """把产物目录指向临时目录。"""
    monkeypatch.setattr(converter, "CONVERT_OUTPUT_DIR", str(tmp_path))
    return tmp_path


# --------------------------------------------------------------------------- #
# 纯文本降级解析
# --------------------------------------------------------------------------- #
def test_parse_plain_table_tab():
    rows = converter._parse_plain_table("A\tB\tC\n1\t2\t3")
    assert rows == [["A", "B", "C"], ["1", "2", "3"]]


def test_parse_plain_table_pipe():
    rows = converter._parse_plain_table("| A | B |\n| --- | --- |\n| 1 | 2 |")
    assert rows == [["A", "B"], ["1", "2"]]


def test_parse_plain_table_returns_none_for_prose():
    assert converter._parse_plain_table("只是一段普通的文字，没有表格。") is None


def test_parse_plain_table_none_for_empty():
    assert converter._parse_plain_table("") is None


# --------------------------------------------------------------------------- #
# CF_HTML / <table> 解析
# --------------------------------------------------------------------------- #
def test_parse_cf_html_extracts_table():
    html = "<html><body>xx<table><tr><td>A</td></tr></table>yy</body></html>"
    extracted = converter._parse_cf_html(html)
    assert extracted.startswith("<table")
    assert extracted.endswith("</table>")


def test_html_parser_simple_table():
    parser = converter.HtmlTableParser()
    parser.feed('<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>')
    assert parser.rows == [["A", "B"], ["1", "2"]]
    assert parser.degraded is False


def test_html_parser_colspan_fills_right():
    parser = converter.HtmlTableParser()
    parser.feed(
        '<table><tr><td colspan="2">A</td><td>B</td></tr>'
        "<tr><td>1</td><td>2</td><td>3</td></tr></table>"
    )
    assert parser.rows == [["A", "A", "B"], ["1", "2", "3"]]


def test_html_parser_rowspan_fills_down():
    parser = converter.HtmlTableParser()
    parser.feed(
        '<table><tr><td rowspan="2">X</td><td>a</td></tr>'
        "<tr><td>b</td></tr></table>"
    )
    assert parser.rows == [["X", "a"], ["", "b"]]


def test_html_parser_nested_table_degrades():
    parser = converter.HtmlTableParser()
    parser.feed(
        "<table><tr><td>outer<table><tr><td>inner</td></tr></table></td></tr></table>"
    )
    assert parser.degraded is True
    assert parser.had_table is True
    assert parser.rows  # 降级但仍产出内容


# --------------------------------------------------------------------------- #
# Markdown → docx
# --------------------------------------------------------------------------- #
_SAMPLE_MD = (
    "# 大标题\n"
    "普通段落 **加粗** 与 *斜体* 与 `code`。\n"
    "- 项目一\n"
    "- 项目二\n"
    "1. 第一\n"
    "2. 第二\n"
    "> 引用内容\n"
    "| 列A | 列B |\n"
    "| --- | --- |\n"
    "| 1 | 2 |\n"
    "```python\n"
    "print(1)\n"
    "```\n"
)


def test_markdown_to_docx_creates_and_readback(outdir):
    result = converter.markdown_to_docx(_SAMPLE_MD)
    assert result.ok is True
    assert os.path.exists(result.path)
    assert result.path.endswith("_md2docx.docx")

    document = docx.Document(result.path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "大标题" in texts
    assert "项目一" in texts
    assert "引用内容" in texts
    assert any(paragraph.style.name.startswith("List") for paragraph in document.paragraphs)
    # 表格：1 行表头 + 1 行数据，2 列。
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 2
    assert len(document.tables[0].columns) == 2


def test_markdown_to_docx_empty(outdir):
    result = converter.markdown_to_docx("")
    assert result.ok is False
    assert result.message == "剪贴板为空"


# --------------------------------------------------------------------------- #
# 剪贴板 → xlsx
# --------------------------------------------------------------------------- #
def test_clipboard_to_xlsx_from_html(outdir, monkeypatch):
    html = (
        "<html><body><table><tr><th>A</th><th>B</th></tr>"
        "<tr><td>1</td><td>2</td></tr></table></body></html>"
    )
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(html=html))
    result = converter.clipboard_to_xlsx()
    assert result.ok is True
    assert result.path.endswith("_html2xlsx.xlsx")

    workbook = openpyxl.load_workbook(result.path)
    sheet = workbook.active
    assert sheet.max_row == 2
    assert sheet.max_column == 2
    assert sheet.cell(1, 1).value == "A"
    assert sheet.cell(2, 2).value == "2"


def test_clipboard_to_xlsx_colspan_alignment(outdir, monkeypatch):
    html = (
        "<table><tr><td colspan='2'>A</td><td>B</td></tr>"
        "<tr><td>1</td><td>2</td><td>3</td></tr></table>"
    )
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(html=html))
    result = converter.clipboard_to_xlsx()
    assert result.ok is True
    workbook = openpyxl.load_workbook(result.path)
    sheet = workbook.active
    assert sheet.max_column == 3
    assert [sheet.cell(1, c).value for c in range(1, 4)] == ["A", "A", "B"]


def test_clipboard_to_xlsx_from_plain_text(outdir, monkeypatch):
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(text="A\tB\n1\t2"))
    result = converter.clipboard_to_xlsx()
    assert result.ok is True
    workbook = openpyxl.load_workbook(result.path)
    assert workbook.active.max_row == 2


def test_clipboard_to_xlsx_empty(outdir, monkeypatch):
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard())
    result = converter.clipboard_to_xlsx()
    assert result.ok is False
    assert result.message == "剪贴板为空"


def test_clipboard_to_xlsx_no_table(outdir, monkeypatch):
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(text="只是一些普通文字。"))
    result = converter.clipboard_to_xlsx()
    assert result.ok is False
    assert result.message == "剪贴板没有可识别的表格"


# --------------------------------------------------------------------------- #
# 打开结果文件夹（不弹窗、不崩溃）
# --------------------------------------------------------------------------- #
def test_open_result_folder_no_crash(outdir, monkeypatch):
    monkeypatch.setattr(converter.os, "startfile", lambda path: None, raising=False)
    converter.open_result_folder()  # 不应抛异常


# --------------------------------------------------------------------------- #
# BUG-1 修复：嵌套表入口提取保留外层闭合标签 → 降级铺格
# --------------------------------------------------------------------------- #
def test_parse_cf_html_keeps_outer_closing_tag():
    nested = "<table><tr><td>x<table><tr><td>y</td></tr></table></td></tr></table>"
    extracted = converter._parse_cf_html(nested)
    assert extracted.count("<table") == 2
    assert extracted.count("</table>") == 2  # 贪婪匹配保留外层闭合标签


def test_clipboard_to_xlsx_nested_table_degrades(outdir, monkeypatch):
    nested = "<table><tr><td>x<table><tr><td>y</td></tr></table></td></tr></table>"
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(html=nested))
    result = converter.clipboard_to_xlsx()
    assert result.ok is True
    assert "铺格" in result.message


def test_clipboard_to_xlsx_html_without_table(outdir, monkeypatch):
    monkeypatch.setattr(
        converter, "clipboard_io", _FakeClipboard(html="<html><body>无表格</body></html>")
    )
    result = converter.clipboard_to_xlsx()
    assert result.ok is False
    assert result.message == "剪贴板没有可识别的表格"  # 有内容但无表格，不得说「剪贴板为空」


# --------------------------------------------------------------------------- #
# OBS-1 修复：同秒内两次转换不覆盖
# --------------------------------------------------------------------------- #
def test_same_second_conversion_does_not_overwrite(outdir, monkeypatch):
    import time as _time

    real_strftime = _time.strftime

    def _fake_strftime(fmt, *args, **kwargs):
        if fmt == converter.CONVERT_FILENAME_PATTERN:
            return "PastePing_20260101_120000"
        return real_strftime(fmt, *args, **kwargs)

    monkeypatch.setattr(converter.time, "strftime", _fake_strftime)
    md = "# 标题\n正文。\n"
    first = converter.markdown_to_docx(md)
    second = converter.markdown_to_docx(md)
    assert first.ok is True and second.ok is True
    assert first.path != second.path
    assert os.path.exists(first.path) and os.path.exists(second.path)
    assert second.path.endswith("_md2docx_1.docx")


# --------------------------------------------------------------------------- #
# 平衡匹配：并列多表 / 嵌套 / 畸形不配平
# --------------------------------------------------------------------------- #
def test_parallel_tables_only_first_parsed(outdir, monkeypatch):
    html = (
        "<table><tr><td>A1</td><td>A2</td></tr></table>"
        "<table><tr><td>B1</td><td>B2</td></tr></table>"
    )
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(html=html))
    result = converter.clipboard_to_xlsx()
    assert result.ok is True

    sheet = openpyxl.load_workbook(result.path).active
    assert sheet.max_row == 1
    assert sheet.max_column == 2
    assert sheet.cell(1, 1).value == "A1"
    assert sheet.cell(1, 2).value == "A2"
    # 第二张表的内容不得混入。
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "B1" not in values and "B2" not in values
    # 明确提示，避免静默丢数据。
    assert "检测到多张表格" in result.message
    assert "第 1 张" in result.message


def test_nested_table_still_degrades(outdir, monkeypatch):
    nested = "<table><tr><td>x<table><tr><td>y</td></tr></table></td></tr></table>"
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(html=nested))
    result = converter.clipboard_to_xlsx()
    assert result.ok is True
    assert "铺格" in result.message
    # 嵌套子表不算独立顶层表，不应误报「多张表格」。
    assert "检测到多张表格" not in result.message


def test_unbalanced_table_no_crash_clear_message(outdir, monkeypatch):
    broken = "<table><tr><td>x</td></tr>"  # 只有 <table>，没有 </table>
    monkeypatch.setattr(converter, "clipboard_io", _FakeClipboard(html=broken))
    result = converter.clipboard_to_xlsx()  # 不得抛异常
    assert result.ok is False
    assert result.message == "剪贴板没有可识别的表格"
