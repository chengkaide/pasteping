"""PastePing 格式转换模块（Markdown → Word、网页表格 → Excel）。

* :func:`markdown_to_docx` —— 解析 Markdown，用 ``python-docx`` 生成 ``.docx``；
* :func:`clipboard_to_xlsx` —— 优先解析剪贴板 ``CF_HTML`` 中的 ``<table>``，
  用 ``openpyxl`` 生成 ``.xlsx``；无 ``CF_HTML`` 时按制表符 / 管道符降级解析；
* :func:`open_result_folder` —— 打开产物目录。

产物自动保存到 ``%USERPROFILE%\\Documents\\PastePing\\``，文件名
``PastePing_YYYYMMDD_HHMMSS_<后缀>.(docx|xlsx)``，**全程零对话框**。

``colspan`` 采用「向右重复填充」、``rowspan`` 采用「向下补空占位」做基础对齐；
无法对齐的复杂嵌套表降级为「按出现顺序铺格」并在通知中附提示，**不报错、不崩溃**。
"""

from __future__ import annotations

import os
import re
import time
import webbrowser
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import Optional

import docx
import openpyxl
from openpyxl.styles import Font

import clipboard_io

# 产物目录：%USERPROFILE%\Documents\PastePing\。
CONVERT_OUTPUT_DIR: str = os.path.join(os.path.expanduser("~"), "Documents", "PastePing")

# 文件名前缀（strftime 模板）。
CONVERT_FILENAME_PATTERN: str = "PastePing_%Y%m%d_%H%M%S"

# 代码块等宽字体。
_MONO_FONT = "Consolas"

# Markdown 行内元素正则（加粗 / 斜体 / 行内代码）。
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*\n]+?\*|_[^_\n]+?_)")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_UNORDERED_RE = re.compile(r"^\s*[-*+]\s+")
_ORDERED_RE = re.compile(r"^\s*\d+[.)]\s+")
_HR_RE = re.compile(r"^[-*_]{3,}$")
_TABLE_SEPARATOR_RE = re.compile(r"^\|?[\s:|-]+\|?$")
# 匹配 <table ...> / </table> 标签（用于平衡匹配第一张表）。
_TABLE_TAG_RE = re.compile(r"</?table\b[^>]*>", re.IGNORECASE)


@dataclass(frozen=True)
class ConvertResult:
    """一次格式转换的结果。

    Attributes:
      ok: 是否成功。
      path: 产物绝对路径；失败时为空串。
      message: 面向用户的通知文案。
    """

    ok: bool
    path: str
    message: str


# --------------------------------------------------------------------------- #
# 通用工具
# --------------------------------------------------------------------------- #
def _collapse(text: str) -> str:
    """折叠空白为单个空格并去首尾。

    Args:
      text: 原始文本。

    Returns:
      折叠后的文本。
    """
    return " ".join(text.split())


def _build_output_path(extension: str, suffix: str) -> str:
    """构造产物路径并确保目录存在；同名文件已存在时追加序号，**绝不覆盖**已有产物。

    Args:
      extension: 扩展名（不含点）。
      suffix: 文件名后缀标签（如 ``_md2docx``）。

    Returns:
      尚不存在的产物绝对路径。
    """
    os.makedirs(CONVERT_OUTPUT_DIR, exist_ok=True)
    stamp = time.strftime(CONVERT_FILENAME_PATTERN)
    base = f"{stamp}{suffix}"
    candidate = os.path.join(CONVERT_OUTPUT_DIR, f"{base}.{extension}")
    index = 1
    while os.path.exists(candidate):
        candidate = os.path.join(CONVERT_OUTPUT_DIR, f"{base}_{index}.{extension}")
        index += 1
    return candidate


# --------------------------------------------------------------------------- #
# Markdown → docx
# --------------------------------------------------------------------------- #
def _add_inline_runs(paragraph, text: str) -> None:
    """把一行 Markdown 行内文本写入段落（支持加粗 / 斜体 / 行内代码）。

    Args:
      paragraph: ``python-docx`` 段落对象。
      text: 行内文本。
    """
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (
            part.startswith("__") and part.endswith("__")
        ):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = _MONO_FONT
        elif (part.startswith("*") and part.endswith("*")) or (
            part.startswith("_") and part.endswith("_")
        ):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _is_table_start(lines: list[str], index: int) -> bool:
    """判断从 ``index`` 起是否为一行 Markdown 管道表格。

    Args:
      lines: 全部行。
      index: 当前行号。

    Returns:
      是表格返回 True。
    """
    line = lines[index]
    if "|" not in line:
        return False
    if index + 1 < len(lines):
        following = lines[index + 1].strip()
        if _TABLE_SEPARATOR_RE.fullmatch(following) and "-" in following and "|" in following:
            return True
    return line.count("|") >= 2


def _parse_pipe_rows(table_lines: list[str]) -> list[list[str]]:
    """把管道表格行解析为二维单元格列表（跳过 ``|---|`` 分隔行）。

    Args:
      table_lines: 表格相关行。

    Returns:
      二维单元格列表。
    """
    rows: list[list[str]] = []
    for line in table_lines:
        stripped = line.strip()
        if _TABLE_SEPARATOR_RE.fullmatch(stripped) and "-" in stripped:
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        rows.append(cells)
    return rows


def _add_table(document, rows: list[list[str]]) -> None:
    """把二维单元格写入 Word 表格。

    Args:
      document: ``python-docx`` 文档对象。
      rows: 二维单元格列表。
    """
    if not rows:
        return
    ncols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for r_index, row in enumerate(rows):
        for c_index in range(ncols):
            cell = table.cell(r_index, c_index)
            cell.text = row[c_index] if c_index < len(row) else ""
            if r_index == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def _add_code_block(document, code_lines: list[str]) -> None:
    """把代码块逐行写入等宽段落。

    Args:
      document: ``python-docx`` 文档对象。
      code_lines: 代码行。
    """
    for line in code_lines:
        run = document.add_paragraph().add_run(line)
        run.font.name = _MONO_FONT


def markdown_to_docx(md_text: Optional[str]) -> ConvertResult:
    """把 Markdown 文本转换为 ``.docx`` 并保存到产物目录。

    Args:
      md_text: Markdown 文本。

    Returns:
      转换结果 ``ConvertResult``。
    """
    if not md_text or not md_text.strip():
        return ConvertResult(ok=False, path="", message="剪贴板为空")

    try:
        document = docx.Document()
        lines = md_text.splitlines()
        index = 0
        total = len(lines)
        while index < total:
            line = lines[index]
            stripped = line.strip()

            # 围栏代码块。
            if stripped.startswith("```"):
                index += 1
                code_lines: list[str] = []
                while index < total and not lines[index].strip().startswith("```"):
                    code_lines.append(lines[index])
                    index += 1
                index += 1  # 跳过结束围栏
                _add_code_block(document, code_lines)
                continue

            # 管道表格。
            if _is_table_start(lines, index):
                table_lines: list[str] = []
                while index < total and "|" in lines[index]:
                    table_lines.append(lines[index])
                    index += 1
                _add_table(document, _parse_pipe_rows(table_lines))
                continue

            # 标题。
            heading = _HEADING_RE.match(stripped)
            if heading:
                level = len(heading.group(1))
                document.add_heading(heading.group(2).strip(), level=min(9, level))
                index += 1
                continue

            # 无序列表。
            if _UNORDERED_RE.match(line):
                paragraph = document.add_paragraph(style="List Bullet")
                _add_inline_runs(paragraph, _UNORDERED_RE.sub("", line, count=1))
                index += 1
                continue

            # 有序列表。
            if _ORDERED_RE.match(line):
                paragraph = document.add_paragraph(style="List Number")
                _add_inline_runs(paragraph, _ORDERED_RE.sub("", line, count=1))
                index += 1
                continue

            # 引用块。
            if stripped.startswith(">"):
                paragraph = document.add_paragraph(style="Intense Quote")
                _add_inline_runs(paragraph, stripped.lstrip(">").strip())
                index += 1
                continue

            # 分隔线 / 空行。
            if _HR_RE.match(stripped) or not stripped:
                index += 1
                continue

            # 普通段落。
            paragraph = document.add_paragraph()
            _add_inline_runs(paragraph, line)
            index += 1

        path = _build_output_path("docx", "_md2docx")
        document.save(path)
        return ConvertResult(ok=True, path=path, message=f"已生成 {os.path.basename(path)}\n{path}")
    except Exception as error:  # noqa: BLE001 - 转换异常须静默降级为失败结果
        return ConvertResult(ok=False, path="", message=f"转换失败：{error}")


# --------------------------------------------------------------------------- #
# CF_HTML 表格解析
# --------------------------------------------------------------------------- #
class HtmlTableParser(HTMLParser):
    """从 HTML 中解析 ``<table>``，产出对齐后的二维行列表。

    Attributes:
      rows: 对齐后的二维单元格列表。
      degraded: 是否因复杂嵌套表而降级（按出现顺序铺格）。
      had_table: 是否解析到 ``<table>``。
    """

    def __init__(self) -> None:
        """初始化解析器。"""
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self.degraded = False
        self.had_table = False
        self._depth = 0
        self._raw_rows: list[list[dict]] = []
        self._current_row: Optional[list[dict]] = None
        self._current_cell: Optional[dict] = None
        self._buffer: list[str] = []
        self._in_cell = False

    @staticmethod
    def _span(value: Optional[str]) -> int:
        """解析 colspan / rowspan 值（非法或缺失时按 1）。

        Args:
          value: 属性原始值。

        Returns:
          跨格数（>=1）。
        """
        try:
            number = int(value) if value is not None else 1
            return number if number >= 1 else 1
        except Exception:
            return 1

    def handle_starttag(self, tag: str, attrs) -> None:
        """处理起始标签。"""
        tag = tag.lower()
        if tag == "table":
            self._depth += 1
            if self._depth > 1:
                self.degraded = True
            else:
                self._raw_rows = []
                self.had_table = True
        elif tag == "tr":
            if self._depth == 1:
                self._current_row = []
                self._raw_rows.append(self._current_row)
        elif tag in ("td", "th"):
            if self._depth == 1:
                attributes = {key.lower(): value for key, value in attrs}
                self._current_cell = {
                    "text": "",
                    "colspan": self._span(attributes.get("colspan")),
                    "rowspan": self._span(attributes.get("rowspan")),
                    "header": tag == "th",
                }
                self._buffer = []
                self._in_cell = True
        elif self._depth > 1:
            self.degraded = True

    def handle_startendtag(self, tag: str, attrs) -> None:
        """处理自闭合标签（``<br/>`` 视为空格）。"""
        if tag.lower() == "br" and self._in_cell:
            self._buffer.append(" ")

    def handle_data(self, data: str) -> None:
        """累积单元格文本。"""
        if self._in_cell:
            self._buffer.append(data)

    def handle_endtag(self, tag: str) -> None:
        """处理结束标签。"""
        tag = tag.lower()
        if tag in ("td", "th"):
            if self._depth == 1 and self._in_cell and self._current_cell is not None:
                self._current_cell["text"] = _collapse("".join(self._buffer))
                if self._current_row is not None:
                    self._current_row.append(self._current_cell)
                self._in_cell = False
                self._current_cell = None
                self._buffer = []
        elif tag == "tr":
            if self._depth == 1:
                self._current_row = None
        elif tag == "table":
            if self._depth == 1:
                self._build_rows()
            self._depth = max(0, self._depth - 1)

    def _build_rows(self) -> None:
        """把原始单元格展开为对齐网格（colspan 向右、rowspan 向下）。"""
        if self.degraded:
            # 复杂嵌套表：降级为「按出现顺序铺格」。
            self.rows = [[cell["text"] for cell in row] for row in self._raw_rows]
            return

        grid: list[list[Optional[str]]] = []
        for r_index, row in enumerate(self._raw_rows):
            while len(grid) <= r_index:
                grid.append([])
            column = 0
            for cell in row:
                while column < len(grid[r_index]) and grid[r_index][column] is not None:
                    column += 1
                colspan = cell["colspan"]
                rowspan = cell["rowspan"]
                for d_row in range(rowspan):
                    target_row = r_index + d_row
                    while len(grid) <= target_row:
                        grid.append([])
                    for d_col in range(colspan):
                        target_col = column + d_col
                        while len(grid[target_row]) <= target_col:
                            grid[target_row].append(None)
                        if d_row == 0:
                            grid[target_row][target_col] = cell["text"]
                        elif grid[target_row][target_col] is None:
                            grid[target_row][target_col] = ""
                column += colspan
        self.rows = [["" if value is None else value for value in row] for row in grid]


def _extract_first_table(raw: str) -> tuple[Optional[str], int]:
    """平衡匹配出 CF_HTML 中的**第一张** ``<table>`` 片段，并统计顶层表数量。

    算法（深度计数）：从第一个 ``<table`` 起扫描标签，``<table`` 让深度 +1、
    ``</table>`` 让深度 −1；**深度首次归零处**即为第一张表的配对闭合标签，在此截断。

    覆盖三种情形：

    1. **嵌套表**：第一张表的配对闭合标签 = 最后一个 ``</table>``（与贪婪一致），
       交给 :class:`HtmlTableParser` 触发 ``degraded`` 降级铺格，不受影响；
    2. **并列多表**（``<table>A</table><table>B</table>``）：只取**第一张**，第二张不并入，
       避免"已闭合又开新表"的畸形结构导致列错位；
    3. **标签不配平**（畸形 HTML，深度永不归零）：降级为"贪婪匹配到最后一个
       ``</table>``"，仍不崩、仍走原有降级 / 失败提示路径。

    顶层表数量 = 深度回到 0 的次数（嵌套子表不计数）。

    Args:
      raw: HTML 文本。

    Returns:
      ``(第一张表片段或 None, 顶层表数量)``。
    """
    depth = 0
    first_start: Optional[int] = None
    first_end: Optional[int] = None
    top_level_count = 0
    for match in _TABLE_TAG_RE.finditer(raw):
        if match.group(0).startswith("</"):
            if depth > 0:
                depth -= 1
                if depth == 0 and first_end is None:
                    first_end = match.end()
        else:
            if depth == 0:
                if first_start is None:
                    first_start = match.start()
                top_level_count += 1
            depth += 1

    if first_start is None:
        return None, 0
    if first_end is None:
        # 情形 3：不配平 → 贪婪到最后一个 </table>（若无则到文本结尾）。
        closes = list(re.finditer(r"</table\s*>", raw, re.IGNORECASE))
        first_end = closes[-1].end() if closes else len(raw)
    return raw[first_start:first_end], top_level_count


def _parse_cf_html(raw: Optional[str]) -> str:
    """从 HTML 文本中抽取**第一张** ``<table>…</table>`` 片段（平衡匹配，见
    :func:`_extract_first_table`）。

    Args:
      raw: HTML 文本（通常来自 ``clipboard_io.read_html`` 的已切片结果）。

    Returns:
      第一张 ``<table>`` 片段；未找到时原样返回。
    """
    if not raw:
        return ""
    fragment, _top_level_count = _extract_first_table(raw)
    return fragment if fragment is not None else raw


def _parse_plain_table(text: str) -> Optional[list[list[str]]]:
    """纯文本降级解析（制表符优先，其次管道符）。

    Args:
      text: 纯文本。

    Returns:
      二维单元格列表；无法识别为表格时返回 None。
    """
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return None

    if any("\t" in line for line in lines):
        rows = [line.split("\t") for line in lines]
        if len({len(row) for row in rows}) == 1 and len(rows[0]) >= 2:
            return [[cell.strip() for cell in row] for row in rows]

    pipe_lines = [line for line in lines if "|" in line]
    if pipe_lines:
        parsed: list[list[str]] = []
        for line in pipe_lines:
            stripped = line.strip()
            if _TABLE_SEPARATOR_RE.fullmatch(stripped) and "-" in stripped:
                continue
            parsed.append([cell.strip() for cell in stripped.strip("|").split("|")])
        if parsed and len({len(row) for row in parsed}) == 1 and len(parsed[0]) >= 2:
            return parsed
    return None


def _write_xlsx(rows: list[list[str]], degrade_note: str) -> ConvertResult:
    """把二维行写入 ``.xlsx`` 并保存。

    Args:
      rows: 二维行列表。
      degrade_note: 降级提示（可为空）。

    Returns:
      转换结果。
    """
    try:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(list(row))
        if sheet.max_row >= 1:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        path = _build_output_path("xlsx", "_html2xlsx")
        workbook.save(path)
        message = f"已生成 {os.path.basename(path)}\n{path}"
        if degrade_note:
            message += f"\n{degrade_note}"
        return ConvertResult(ok=True, path=path, message=message)
    except Exception as error:  # noqa: BLE001
        return ConvertResult(ok=False, path="", message=f"转换失败：{error}")


def clipboard_to_xlsx() -> ConvertResult:
    """读取剪贴板并转换为 ``.xlsx``。

    优先解析 ``CF_HTML`` 中的 ``<table>``（含简单跨格对齐；复杂嵌套表降级铺格）；
    无 ``CF_HTML`` 时按制表符 / 管道符降级解析。

    失败文案区分两种原因：剪贴板**确实为空/纯空白** → 「剪贴板为空」；
    剪贴板**有内容但未解析出表格** → 「剪贴板没有可识别的表格」（PRD CNV-6）。

    Returns:
      转换结果 ``ConvertResult``。
    """
    try:
        html = clipboard_io.read_html()
    except Exception:
        html = None
    try:
        text = clipboard_io.read_text()
    except Exception:
        text = None

    has_text = bool(text and text.strip())

    if html:
        fragment, table_count = _extract_first_table(html)
        parser = HtmlTableParser()
        try:
            parser.feed(fragment if fragment is not None else html)
            parser.close()
        except Exception:
            parser.rows = []
        if parser.rows and any(any(cell for cell in row) for row in parser.rows):
            notes: list[str] = []
            if parser.degraded:
                notes.append("表格结构较复杂，已按顺序铺格")
            if table_count > 1:
                # 只转换了第一张表，明确提示，避免静默丢数据。
                notes.append("检测到多张表格，已转换第 1 张")
            return _write_xlsx(parser.rows, "\n".join(notes))

    if has_text:
        rows = _parse_plain_table(text)
        if rows:
            return _write_xlsx(rows, "")
        return ConvertResult(ok=False, path="", message="剪贴板没有可识别的表格")

    # 剪贴板连文本 / CF_HTML 都没有：才是真正的「空」。
    if not html:
        return ConvertResult(ok=False, path="", message="剪贴板为空")
    # 有 CF_HTML 但未解析出表格：属于「有内容、无表格」。
    return ConvertResult(ok=False, path="", message="剪贴板没有可识别的表格")


def open_result_folder() -> None:
    """打开产物目录（失败静默）。"""
    try:
        os.makedirs(CONVERT_OUTPUT_DIR, exist_ok=True)
        os.startfile(CONVERT_OUTPUT_DIR)  # type: ignore[attr-defined]  # 仅 Windows
    except Exception:
        try:
            webbrowser.open(CONVERT_OUTPUT_DIR)
        except Exception:
            pass
